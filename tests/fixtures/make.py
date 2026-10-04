"""Synthetic fixture builders for the identify/ extractor tests.

Generated at test time rather than committed as binary blobs: deterministic,
easy to tweak per test, and avoids checking in things like a real ROM dump
(which wouldn't be ours to redistribute anyway).
"""

from __future__ import annotations

import io
import struct
import tarfile
import wave
import zipfile
from pathlib import Path

SECTOR = 2048


def make_zip(path: Path, members: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)


def make_tar_gz(path: Path, members: dict[str, bytes]) -> None:
    with tarfile.open(path, "w:gz") as tf:
        for name, data in members.items():
            info = tarfile.TarInfo(name)
            info.size = len(data)
            tf.addfile(info, io.BytesIO(data))


def make_wav(path: Path, seconds: float = 1.0, framerate: int = 8000) -> None:
    with wave.open(str(path), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(framerate)
        w.writeframes(b"\x00\x00" * int(seconds * framerate))


def make_png(
    path: Path, size: tuple[int, int] = (1920, 1080), exif: dict[str, str] | None = None
) -> None:
    from PIL import ExifTags, Image

    img = Image.new("RGB", size, (10, 20, 30))
    if exif:
        name_to_id = {v: k for k, v in ExifTags.TAGS.items()}
        exif_obj = img.getexif()
        for field, value in exif.items():
            exif_obj[name_to_id[field]] = value
        img.save(path, exif=exif_obj.tobytes())
    else:
        img.save(path)


def make_pdf(path: Path, text: str = "Hello World") -> None:
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
    ]
    stream = f"BT /F1 24 Tf 10 100 Td ({text}) Tj ET".encode("latin-1")
    objs.append(b"<</Length %d>>stream\n%s\nendstream" % (len(stream), stream))
    objs.append(b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj".encode() + b"\n" + body + b"\nendobj\n"

    xref_offset = len(out)
    n = len(objs) + 1
    out += f"xref\n0 {n}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<</Size {n}/Root 1 0 R>>\nstartxref\n{xref_offset}\n%%EOF".encode()
    path.write_bytes(bytes(out))


def make_docx(path: Path, title: str = "Test Doc", text: str = "Hello from docx") -> None:
    import docx

    document = docx.Document()
    document.core_properties.title = title
    document.add_paragraph(text)
    document.save(str(path))


def make_xlsx(path: Path, sheet_names: list[str] = ("Sheet1", "Data")) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.title = sheet_names[0]
    for name in sheet_names[1:]:
        workbook.create_sheet(name)
    workbook.save(str(path))


def make_pptx(path: Path, title: str = "My Slide Title") -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = title
    presentation.save(str(path))


def make_font(path: Path, family: str = "Test Sans", style: str = "Regular") -> None:
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.ttGlyphPen import TTGlyphPen

    fb = FontBuilder(1000, isTTF=True)
    fb.setupGlyphOrder([".notdef", "A"])
    fb.setupCharacterMap({65: "A"})
    pen = TTGlyphPen(None)
    pen.moveTo((0, 0))
    pen.lineTo((0, 500))
    pen.lineTo((500, 500))
    pen.closePath()
    glyph = pen.glyph()
    fb.setupGlyf({".notdef": glyph, "A": glyph})
    fb.setupHorizontalMetrics({".notdef": (500, 0), "A": (500, 0)})
    fb.setupHorizontalHeader(ascent=800, descent=-200)
    fb.setupNameTable({"familyName": family, "styleName": style})
    fb.setupOS2()
    fb.setupPost()
    fb.save(str(path))


def _bencode(value) -> bytes:
    if isinstance(value, int):
        return f"i{value}e".encode()
    if isinstance(value, (bytes, str)):
        raw = value.encode() if isinstance(value, str) else value
        return str(len(raw)).encode() + b":" + raw
    if isinstance(value, dict):
        body = b"".join(_bencode(k) + _bencode(v) for k in sorted(value) for v in (value[k],))
        return b"d" + body + b"e"
    raise TypeError(value)


def make_torrent(path: Path, name: str = "test.iso", length: int = 12345) -> None:
    path.write_bytes(_bencode({"info": {"name": name, "length": length}}))


def _dir_record(name: bytes, extent: int, length: int, is_dir: bool) -> bytes:
    flags = 2 if is_dir else 0
    rec = bytearray(
        [0, 0]
        + list(struct.pack("<I", extent) + struct.pack(">I", extent))
        + list(struct.pack("<I", length) + struct.pack(">I", length))
        + [0] * 7
        + [flags, 0, 0]
        + list(struct.pack("<H", 1) + struct.pack(">H", 1))
        + [len(name)]
        + list(name)
    )
    if len(name) % 2 == 0:
        rec.append(0)
    rec[0] = len(rec)
    return bytes(rec)


def make_iso9660(
    path: Path, volume_id: str = "", root_files: dict[str, bytes] | None = None
) -> None:
    """A minimal, plain (2048-byte sector) ISO9660 image: PVD + root dir + files."""
    root_files = root_files or {}
    data_sector = 19
    layout: dict[str, tuple[int, int, bytes]] = {}
    cursor = data_sector
    for name, content in root_files.items():
        layout[name] = (cursor, len(content), content)
        cursor += max(1, (len(content) + SECTOR - 1) // SECTOR)

    root_entries = bytearray()
    root_entries += _dir_record(b"\x00", 18, SECTOR, True)
    root_entries += _dir_record(b"\x01", 18, SECTOR, True)
    for name, (extent, length, _content) in layout.items():
        root_entries += _dir_record(f"{name};1".encode("ascii"), extent, length, False)

    image = bytearray(cursor * SECTOR)

    pvd = bytearray(SECTOR)
    pvd[0] = 1
    pvd[1:6] = b"CD001"
    pvd[6] = 1
    pvd[40:72] = volume_id.encode("ascii").ljust(32)[:32]
    root_record = _dir_record(b"\x00", 18, SECTOR, True)
    pvd[156 : 156 + len(root_record)] = root_record
    image[16 * SECTOR : 17 * SECTOR] = pvd

    terminator = bytearray(SECTOR)
    terminator[0] = 255
    terminator[1:6] = b"CD001"
    image[17 * SECTOR : 18 * SECTOR] = terminator

    image[18 * SECTOR : 18 * SECTOR + len(root_entries)] = root_entries
    for _name, (extent, length, content) in layout.items():
        image[extent * SECTOR : extent * SECTOR + length] = content

    path.write_bytes(bytes(image))


def make_disc_magic(path: Path, offset: int, magic: bytes, total_size: int = SECTOR * 2) -> None:
    """A blank image with just a console disc magic at a fixed offset (GC/Wii)."""
    data = bytearray(total_size)
    data[offset : offset + len(magic)] = magic
    path.write_bytes(bytes(data))
