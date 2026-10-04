"""Office document core properties and a content preview.
See PLAN.md section 4.3.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import docx
import openpyxl
from pptx import Presentation
from pptx.exc import PackageNotFoundError

from .base import Extractor
from .types import Evidence

PREVIEW_LINES = 40


class OfficeExtractor(Extractor):
    name = "office"
    priority = 36

    def can_handle(self, evidence: Evidence) -> bool:
        return evidence.extension in (".docx", ".xlsx", ".pptx")

    def extract(self, path: Path, evidence: Evidence) -> dict[str, Any]:
        try:
            if evidence.extension == ".docx":
                return self._docx(path)
            if evidence.extension == ".xlsx":
                return self._xlsx(path)
            return self._pptx(path)
        except (OSError, KeyError, PackageNotFoundError):
            return {}

    def _core_props(self, props) -> dict[str, Any]:
        result: dict[str, Any] = {}
        if getattr(props, "title", None):
            result["title"] = props.title
        # openpyxl calls this `creator`; python-docx/python-pptx call it `author`.
        author = getattr(props, "author", None) or getattr(props, "creator", None)
        if author:
            result["author"] = author
        return result

    def _docx(self, path: Path) -> dict[str, Any]:
        document = docx.Document(str(path))
        result = self._core_props(document.core_properties)
        lines = [p.text for p in document.paragraphs if p.text.strip()][:PREVIEW_LINES]
        result["text_preview"] = "\n".join(lines)
        return result

    def _xlsx(self, path: Path) -> dict[str, Any]:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            result = self._core_props(workbook.properties)
            result["sheet_names"] = workbook.sheetnames
            return result
        finally:
            workbook.close()

    def _pptx(self, path: Path) -> dict[str, Any]:
        presentation = Presentation(str(path))
        result = self._core_props(presentation.core_properties)
        titles = []
        for slide in presentation.slides:
            if slide.shapes.title and slide.shapes.title.text:
                titles.append(slide.shapes.title.text)
        result["slide_titles"] = titles[:PREVIEW_LINES]
        return result
